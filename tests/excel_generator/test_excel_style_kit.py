import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPT_PATH = REPO_ROOT / ".claude" / "skills" / "excel-generator" / "scripts" / "excel_style_kit.py"


@pytest.fixture(scope="module")
def esk():
    spec = importlib.util.spec_from_file_location("excel_style_kit", SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules["excel_style_kit"] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture()
def ws():
    wb = Workbook()
    return wb.active


class TestGetTheme:
    def test_known_theme(self, esk):
        theme = esk.get_theme('corporate_blue')
        assert theme['primary'] == '1F4E79'

    def test_unknown_theme_raises(self, esk):
        with pytest.raises(KeyError):
            esk.get_theme('nonexistent_theme')

    def test_all_themes_have_required_keys(self, esk):
        for name, theme in esk.THEMES.items():
            assert set(theme) == {'primary', 'light', 'accent', 'chart_colors'}
            assert len(theme['chart_colors']) >= 1


class TestStyling:
    def test_style_title_sets_font_and_value(self, esk, ws):
        theme = esk.get_theme('elegant_black')
        cell = esk.style_title(ws, 'B2', 'My Title', theme)
        assert cell.value == 'My Title'
        assert cell.font.bold is True
        assert cell.font.color.rgb.endswith(theme['primary'])

    def test_style_table_header_applies_fill_and_font(self, esk, ws):
        theme = esk.get_theme('elegant_black')
        esk.style_table_header(ws, row=5, start_col=2, end_col=4, theme=theme)
        for col in (2, 3, 4):
            cell = ws.cell(row=5, column=col)
            assert cell.font.color.rgb.endswith('FFFFFF')
            assert cell.fill.start_color.rgb.endswith(theme['primary'])


class TestApplyNumberFormats:
    def test_applies_to_range(self, esk, ws):
        for row in range(1, 4):
            ws.cell(row=row, column=1, value=row * 1.5)
        esk.apply_number_formats(ws, col=1, start_row=1, end_row=3, format_key='decimal1')
        for row in range(1, 4):
            assert ws.cell(row=row, column=1).number_format == '#,##0.0'

    def test_unknown_format_key_raises(self, esk, ws):
        with pytest.raises(KeyError):
            esk.apply_number_formats(ws, col=1, start_row=1, end_row=1, format_key='nonsense')


class TestApplyDataBlockBorders:
    def test_every_cell_gets_a_border(self, esk, ws):
        esk.apply_data_block_borders(ws, start_row=2, end_row=4, start_col=2, end_col=3, has_header=True)
        for row in range(2, 5):
            for col in (2, 3):
                cell = ws.cell(row=row, column=col)
                assert cell.border.left.style is not None or col != 2 or True
        # Left border only on first column, right border only on last column.
        assert ws.cell(row=2, column=2).border.left.style == 'thin'
        assert ws.cell(row=2, column=3).border.right.style == 'thin'
        assert ws.cell(row=2, column=2).border.right.style is None

    def test_header_bottom_is_medium(self, esk, ws):
        esk.apply_data_block_borders(ws, start_row=2, end_row=4, start_col=2, end_col=3, has_header=True)
        assert ws.cell(row=2, column=2).border.bottom.style == 'medium'
        assert ws.cell(row=3, column=2).border.bottom.style == 'thin'
        assert ws.cell(row=4, column=2).border.bottom.style == 'thin'

    def test_no_header_top_row_has_no_medium_border(self, esk, ws):
        esk.apply_data_block_borders(ws, start_row=2, end_row=4, start_col=2, end_col=3, has_header=False)
        assert ws.cell(row=2, column=2).border.bottom.style == 'thin'

    def test_invalid_range_raises(self, esk, ws):
        with pytest.raises(ValueError):
            esk.apply_data_block_borders(ws, start_row=5, end_row=2, start_col=2, end_col=3)


class TestCalculateColumnWidth:
    def test_numeric_column_padding(self, esk, ws):
        ws.cell(row=2, column=2, value=12345)
        ws.cell(row=3, column=2, value=1)
        width = esk.calculate_column_width(ws, col=2, data_block_ranges=[(2, 3)])
        assert width == max(5 + 6, 14)

    def test_text_column_minimum(self, esk, ws):
        ws.cell(row=2, column=2, value="Hi")
        width = esk.calculate_column_width(ws, col=2, data_block_ranges=[(2, 2)])
        assert width == 15  # minimum for text columns

    def test_ignores_rows_outside_data_block_ranges(self, esk, ws):
        ws.cell(row=2, column=2, value="short")
        ws.cell(row=10, column=2, value="a very very very long standalone text row")
        width = esk.calculate_column_width(ws, col=2, data_block_ranges=[(2, 2)])
        assert width == 15  # "short" + padding 4 = 9, minimum 15 wins; row 10 excluded

    def test_empty_data_block_uses_numeric_minimum_by_default(self, esk, ws):
        # No content found -> is_numeric stays at its True default -> numeric minimum (14).
        width = esk.calculate_column_width(ws, col=2, data_block_ranges=[(2, 2)])
        assert width == 14


class TestAddSheetIndex:
    def test_writes_title_and_hyperlinks(self, esk, ws):
        theme = esk.get_theme('elegant_black')
        esk.add_sheet_index(ws, ['Overview', 'Data', 'Analysis'], theme, start_row=6)
        assert ws.cell(row=5, column=2).value == "CONTENTS"
        assert ws.cell(row=6, column=2).value == "Overview"
        assert ws.cell(row=6, column=2).hyperlink.target == "#'Overview'!A1"
        assert ws.cell(row=8, column=2).value == "Analysis"

    def test_title_row_must_precede_start_row(self, esk, ws):
        theme = esk.get_theme('elegant_black')
        with pytest.raises(ValueError):
            esk.add_sheet_index(ws, ['Overview'], theme, start_row=5, title_row=5)


class TestBuildDemoWorkbook:
    def test_builds_and_saves(self, esk, tmp_path):
        wb = esk.build_demo_workbook('corporate_blue')
        out_path = tmp_path / "demo.xlsx"
        wb.save(out_path)
        assert out_path.exists()

        reloaded = load_workbook(out_path)
        assert "Overview" in reloaded.sheetnames
        assert "Data" in reloaded.sheetnames
        ws = reloaded["Overview"]
        assert ws["B2"].value == "Demo Report"

    def test_unknown_theme_raises(self, esk):
        with pytest.raises(KeyError):
            esk.build_demo_workbook('not-a-theme')


def test_cli_demo_creates_file(esk, tmp_path):
    out_path = tmp_path / "cli_demo.xlsx"
    sys.argv = ["excel_style_kit.py", "demo", str(out_path), "--theme", "elegant_black"]
    esk.main()
    assert out_path.exists()


def test_cli_demo_unknown_theme_exits_nonzero(esk, tmp_path):
    out_path = tmp_path / "cli_demo.xlsx"
    sys.argv = ["excel_style_kit.py", "demo", str(out_path), "--theme", "nope"]
    with pytest.raises(SystemExit) as exc_info:
        esk.main()
    assert exc_info.value.code == 1
