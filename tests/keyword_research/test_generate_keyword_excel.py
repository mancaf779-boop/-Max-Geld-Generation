"""End-to-end and unit tests for scripts/generate_keyword_excel.py."""
import json
import subprocess
import sys

import openpyxl
import pytest

from . import conftest

SCRIPT = conftest.SKILL_SCRIPTS / "generate_keyword_excel.py"
FIXTURES = conftest.FIXTURES

sys.path.insert(0, str(conftest.SKILL_SCRIPTS))
import generate_keyword_excel as gke  # noqa: E402


def run_script(args, cwd=None):
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        cwd=cwd, capture_output=True, text=True,
    )


# ------------------------- CLI / usage --------------------------------------
def test_help_exits_zero():
    result = run_script(["--help"])
    assert result.returncode == 0
    assert "usage" in result.stdout.lower()


def test_no_args_exits_nonzero_with_usage():
    result = run_script([])
    assert result.returncode != 0


def test_missing_input_file_exits_cleanly(tmp_path):
    out = tmp_path / "out.xlsx"
    result = run_script([str(tmp_path / "does_not_exist.json"), str(out)])
    assert result.returncode != 0
    assert "not found" in (result.stdout + result.stderr).lower()
    assert not out.exists()


def test_malformed_json_exits_cleanly(tmp_path):
    bad = tmp_path / "bad.json"
    bad.write_text("{not valid json")
    out = tmp_path / "out.xlsx"
    result = run_script([str(bad), str(out)])
    assert result.returncode != 0
    assert "not valid json" in (result.stdout + result.stderr).lower()


def test_json_not_an_object_exits_cleanly(tmp_path):
    bad = tmp_path / "list.json"
    bad.write_text("[1, 2, 3]")
    out = tmp_path / "out.xlsx"
    result = run_script([str(bad), str(out)])
    assert result.returncode != 0
    assert "must be an object" in (result.stdout + result.stderr).lower()


# ------------------------- end-to-end workbook structure ---------------------
def test_small_dataset_generates_workbook_with_warning(tmp_path):
    out = tmp_path / "keyword_research.xlsx"
    result = run_script([str(FIXTURES / "small_keyword_data.json"), str(out)])
    assert result.returncode == 0, result.stderr
    assert "Expected 300 keywords, got 6" in result.stdout
    assert out.exists()

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Topic Cluster Summary", "Priority Keyword Targets", "Competitor Gaps & Landscape"]

    ws_kw = wb["Priority Keyword Targets"]
    header = [c.value for c in ws_kw[1]]
    assert header == [
        "Keyword", "Topic Cluster", "Search Intent", "Monthly Volume (US)",
        "Global Volume", "Keyword Difficulty (KD)", "KD Level", "CPC",
        "Parent Topic", "Recommended Content Type", "Priority",
    ]
    # 6 data rows + 1 header row.
    assert ws_kw.max_row == 7

    ws_summary = wb["Topic Cluster Summary"]
    assert ws_summary["A1"].value == "Topic Cluster Search Volume Summary"


def test_full_300_keyword_dataset(full_payload_path, tmp_path):
    out = tmp_path / "keyword_research.xlsx"
    result = run_script([str(full_payload_path), str(out)])
    assert result.returncode == 0, result.stderr
    assert "Expected 300 keywords" not in result.stdout

    wb = openpyxl.load_workbook(out)
    ws_kw = wb["Priority Keyword Targets"]
    # header + 300 data rows
    assert ws_kw.max_row == 301

    ws_summary = wb["Topic Cluster Summary"]
    # Subtitle row states 300 keywords across 5 clusters.
    assert "300 Keywords" in ws_summary["A2"].value
    assert "5 Topic Clusters" in ws_summary["A2"].value

    ws_comp = wb["Competitor Gaps & Landscape"]
    values = [ws_comp.cell(row=r, column=1).value for r in range(1, ws_comp.max_row + 1)]
    assert "Section 1: Traffic Share by Domain" in values
    assert "Section 2: Keyword Gap Opportunities" in values
    assert "Section 3: Strategic Insights & Recommendations" in values


def test_priority_color_coding(tmp_path):
    out = tmp_path / "keyword_research.xlsx"
    run_script([str(FIXTURES / "small_keyword_data.json"), str(out)])
    wb = openpyxl.load_workbook(out)
    ws = wb["Priority Keyword Targets"]
    header = [c.value for c in ws[1]]
    priority_idx = header.index("Priority") + 1
    # Row 2 = "best ai writer" -> High priority -> green fill.
    fill = ws.cell(row=2, column=priority_idx).fill
    assert fill.start_color.rgb in ("00C6EFCE", "FFC6EFCE", "C6EFCE")


# ------------------------- volume-column auto-detection -----------------------
def test_uk_volume_column_is_detected_and_used(tmp_path):
    payload = conftest.make_full_payload(30, volume_col="Monthly Volume (UK)")
    path = tmp_path / "uk_data.json"
    path.write_text(json.dumps(payload))
    out = tmp_path / "out.xlsx"

    result = run_script([str(path), str(out)])
    assert result.returncode == 0, result.stderr

    wb = openpyxl.load_workbook(out)
    ws_kw = wb["Priority Keyword Targets"]
    header = [c.value for c in ws_kw[1]]
    assert "Monthly Volume (UK)" in header
    assert "Monthly Volume (US)" not in header

    ws_summary = wb["Topic Cluster Summary"]
    summary_header = [ws_summary.cell(row=4, column=c).value for c in range(1, 10)]
    assert "Monthly Volume (UK)" in summary_header


def test_explicit_volume_col_override(tmp_path):
    payload = conftest.make_full_payload(10, volume_col="Monthly Volume (DE)")
    path = tmp_path / "de_data.json"
    path.write_text(json.dumps(payload))
    out = tmp_path / "out.xlsx"

    result = run_script([str(path), str(out), "--volume-col", "Monthly Volume (DE)"])
    assert result.returncode == 0, result.stderr
    wb = openpyxl.load_workbook(out)
    header = [c.value for c in wb["Priority Keyword Targets"][1]]
    assert "Monthly Volume (DE)" in header


def test_global_volume_only_scope_falls_back_gracefully(tmp_path):
    # Worldwide scope: keywords have no "Monthly Volume (...)" key at all.
    keywords = conftest.make_keywords(10)
    for kw in keywords:
        kw.pop("Monthly Volume (US)", None)
    payload = {
        "keywords": keywords, "competitors": [], "gaps": [],
        "insights": {"tam": "", "opportunities": "", "findings": [], "recommendations": []},
    }
    path = tmp_path / "global_only.json"
    path.write_text(json.dumps(payload))
    out = tmp_path / "out.xlsx"

    result = run_script([str(path), str(out)])
    assert result.returncode == 0, result.stderr
    wb = openpyxl.load_workbook(out)
    # Falls back to the documented default label; column exists (blank) rather than crashing.
    header = [c.value for c in wb["Priority Keyword Targets"][1]]
    assert "Monthly Volume (US)" in header


# ------------------------- unit tests for pure helpers ------------------------
def test_num_helper_handles_none_and_blank_and_bad_values():
    assert gke._num(None) == 0
    assert gke._num("") == 0
    assert gke._num("not a number") == 0
    assert gke._num("42") == 42.0
    assert gke._num(3.5) == 3.5
    assert gke._num(None, default=-1) == -1


def test_detect_volume_col_prefers_override():
    keywords = [{"Monthly Volume (US)": 100}]
    assert gke.detect_volume_col(keywords, override="Monthly Volume (FR)") == "Monthly Volume (FR)"


def test_detect_volume_col_scans_keys():
    keywords = [{"Global Volume": 100}, {"Monthly Volume (JP)": 50}]
    assert gke.detect_volume_col(keywords) == "Monthly Volume (JP)"


def test_detect_volume_col_defaults_when_absent():
    assert gke.detect_volume_col([]) == "Monthly Volume (US)"


def test_create_cluster_summary_handles_none_values_without_crash():
    """Keywords with null KD/CPC/volume (valid JSON, e.g. worldwide scope
    partially filled) must not raise TypeError during aggregation."""
    from openpyxl import Workbook
    wb = Workbook()
    keywords = [
        {"Topic Cluster": "A", "Monthly Volume (US)": None, "Global Volume": None,
         "Keyword Difficulty (KD)": None, "CPC": None, "Priority": "High"},
        {"Topic Cluster": "A", "Monthly Volume (US)": 500, "Global Volume": 700,
         "Keyword Difficulty (KD)": 20, "CPC": 1.5, "Priority": "Low"},
    ]
    ws = gke.create_cluster_summary(wb, keywords)
    assert ws["A1"].value == "Topic Cluster Search Volume Summary"


def test_create_cluster_summary_empty_keywords_no_crash():
    from openpyxl import Workbook
    wb = Workbook()
    ws = gke.create_cluster_summary(wb, [])
    assert ws is not None
