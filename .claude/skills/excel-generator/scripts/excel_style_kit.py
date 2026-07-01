#!/usr/bin/env python3
"""
Excel styling helper library for the excel-generator skill.

Consolidates the theme system, data-block borders, column-width calculation,
and sheet-index patterns documented in SKILL.md into reusable, tested
functions built on openpyxl.

CLI usage (for a quick visual sanity check of a theme):
    excel_style_kit.py demo <output.xlsx> [--theme elegant_black]
"""

import argparse
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

SERIF_FONT = 'Georgia'
SANS_FONT = 'Calibri'

THEMES = {
    'elegant_black': {
        'primary': '2D2D2D', 'light': 'E5E5E5', 'accent': '2D2D2D',
        'chart_colors': ['2D2D2D', '4A4A4A', '6B6B6B', '8C8C8C', 'ADADAD', 'CFCFCF'],
    },
    'corporate_blue': {
        'primary': '1F4E79', 'light': 'D6E3F0', 'accent': '1F4E79',
        'chart_colors': ['1F4E79', '2E75B6', '5B9BD5', '9DC3E6', 'BDD7EE', 'DEEBF7'],
    },
    'forest_green': {
        'primary': '2E5A4C', 'light': 'D4E5DE', 'accent': '2E5A4C',
        'chart_colors': ['2E5A4C', '3F7A67', '5A9B85', '85BCAB', 'ADD4C7', 'D4E5DE'],
    },
    'burgundy': {
        'primary': '722F37', 'light': 'E8D5D7', 'accent': '722F37',
        'chart_colors': ['722F37', '954049', 'B8626B', 'D48F95', 'E4B8BC', 'F0DADD'],
    },
    'slate_gray': {
        'primary': '4A5568', 'light': 'E2E8F0', 'accent': '4A5568',
        'chart_colors': ['4A5568', '667080', '8896A3', 'AAB4BE', 'CCD3D9', 'E2E8F0'],
    },
}

SEMANTIC_COLORS = {
    'positive': '2E7D32',
    'negative': 'C62828',
    'warning': 'F57C00',
}

HIGHLIGHT_COLORS = {
    'emphasis': 'E6F3FF',
    'section': 'FFF3E0',
    'input': 'FFFDE7',
    'special': 'FFF9C4',
    'success': 'E8F5E9',
    'warning': 'FFCCBC',
}

NUMBER_FORMATS = {
    'integer': '#,##0',
    'decimal1': '#,##0.0',
    'decimal2': '#,##0.00',
    'percentage': '0.0%',
    'currency': '$#,##0.00',
}


def get_theme(name):
    """Look up a theme by name, raising a clear error for unknown themes."""
    if name not in THEMES:
        raise KeyError(f"Unknown theme '{name}'. Available: {', '.join(sorted(THEMES))}")
    return THEMES[name]


def style_title(ws: Worksheet, cell_ref, text, theme):
    cell = ws[cell_ref]
    cell.value = text
    cell.font = Font(name=SERIF_FONT, size=18, bold=True, color=theme['primary'])
    return cell


def style_section_header(ws: Worksheet, cell_ref, text, theme, with_fill=False):
    cell = ws[cell_ref]
    cell.value = text
    cell.font = Font(name=SERIF_FONT, size=14, bold=True, color=theme['primary'])
    if with_fill:
        cell.fill = PatternFill(start_color=theme['light'], end_color=theme['light'], fill_type='solid')
    return cell


def style_table_header(ws: Worksheet, row, start_col, end_col, theme):
    """Apply header styling (white bold serif text on theme-primary fill) to a header row."""
    header_font = Font(name=SERIF_FONT, size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color=theme['primary'], end_color=theme['primary'], fill_type='solid')
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')


def apply_number_formats(ws: Worksheet, col, start_row, end_row, format_key):
    """Apply a named number format to every cell in a column range (values and formulas alike)."""
    if format_key not in NUMBER_FORMATS:
        raise KeyError(f"Unknown format_key '{format_key}'. Available: {', '.join(sorted(NUMBER_FORMATS))}")
    fmt = NUMBER_FORMATS[format_key]
    for row in range(start_row, end_row + 1):
        ws.cell(row=row, column=col).number_format = fmt


def apply_data_block_borders(ws: Worksheet, start_row, end_row, start_col, end_col, has_header=True):
    """
    Apply borders to a Data Block: outer frame on all 4 sides, medium header-bottom
    border if has_header, thin internal horizontal borders between rows, no
    internal vertical borders. Every cell in the range is processed.
    """
    if end_row < start_row or end_col < start_col:
        raise ValueError("end_row/end_col must be >= start_row/start_col")

    outer_border = Side(style='thin', color='D1D1D1')
    header_bottom = Side(style='medium', color='2D2D2D')
    inner_horizontal = Side(style='thin', color='D1D1D1')
    no_border = Side(style=None)

    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)

            left = outer_border if col == start_col else no_border
            right = outer_border if col == end_col else no_border
            top = outer_border if row == start_row else inner_horizontal

            if has_header and row == start_row:
                bottom = header_bottom
            elif row == end_row:
                bottom = outer_border
            else:
                bottom = inner_horizontal

            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def calculate_column_width(ws: Worksheet, col, data_block_ranges):
    """
    Calculate a column's width based only on Data Block content
    (standalone text rows, section headers, and notes are excluded by
    virtue of not being passed in data_block_ranges).
    """
    max_len = 0
    is_numeric = True

    for start_row, end_row in data_block_ranges:
        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value not in (None, ''):
                display_value = str(cell.value)
                max_len = max(max_len, len(display_value))
                if not isinstance(cell.value, (int, float)):
                    is_numeric = False

    padding = 6 if is_numeric else 4
    minimum = 14 if is_numeric else 15
    return max(max_len + padding, minimum)


def add_sheet_index(ws: Worksheet, sheet_names, theme, start_row=6, col=2, title_row=None):
    """Write a CONTENTS section with hyperlinks to each sheet, starting at start_row."""
    if title_row is None:
        title_row = start_row - 1
    if title_row >= start_row:
        raise ValueError("title_row must be before start_row")

    title_cell = ws.cell(row=title_row, column=col, value="CONTENTS")
    title_cell.font = Font(name=SERIF_FONT, size=14, bold=True, color=theme['accent'])

    for i, sheet_name in enumerate(sheet_names, start=start_row):
        cell = ws.cell(row=i, column=col, value=sheet_name)
        cell.hyperlink = f"#'{sheet_name}'!A1"
        cell.font = Font(color=theme['accent'], underline='single')


def build_demo_workbook(theme_name='elegant_black'):
    """Build a small workbook exercising the theme, data-block borders, and sheet index."""
    theme = get_theme(theme_name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3

    style_title(ws, 'B2', 'Demo Report', theme)
    add_sheet_index(ws, ['Overview', 'Data'], theme, start_row=6)

    style_section_header(ws, 'B10', 'KEY METRICS', theme, with_fill=True)
    headers = ['Metric', 'Value']
    for i, h in enumerate(headers, start=2):
        ws.cell(row=12, column=i, value=h)
    ws.cell(row=13, column=2, value='Revenue')
    ws.cell(row=13, column=3, value=1234567)
    ws.cell(row=14, column=2, value='Growth')
    ws.cell(row=14, column=3, value=0.23)

    style_table_header(ws, row=12, start_col=2, end_col=3, theme=theme)
    apply_data_block_borders(ws, start_row=12, end_row=14, start_col=2, end_col=3, has_header=True)
    apply_number_formats(ws, col=3, start_row=13, end_row=13, format_key='integer')
    apply_number_formats(ws, col=3, start_row=14, end_row=14, format_key='percentage')

    width = calculate_column_width(ws, col=2, data_block_ranges=[(12, 14)])
    ws.column_dimensions['B'].width = width

    wb.create_sheet('Data')
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    p_demo = sub.add_parser("demo", help="Generate a small sample workbook")
    p_demo.add_argument("output_path")
    p_demo.add_argument("--theme", default="elegant_black")

    args = parser.parse_args()

    if args.command == "demo":
        try:
            wb = build_demo_workbook(args.theme)
        except KeyError as e:
            print(f"Error: {e}", file=sys.stderr)
            sys.exit(1)
        wb.save(args.output_path)
        print(f"Wrote demo workbook to {args.output_path}")


if __name__ == "__main__":
    main()
